import openpyxl

wb = openpyxl.load_workbook(r'11_Demos_and_Workbooks\05_Pivot_Tables\Module_5_Demo.xlsx', read_only=True, data_only=True)
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n================ SHEET: {sheetname} ================")
    count = 0
    for row in ws.iter_rows(values_only=True):
        count += 1
        if count <= 15:
            non_empty = [str(x) for x in row if x is not None]
            if non_empty:
                print(f"Row {count}: {non_empty[:10]}")
    print(f"Total rows in {sheetname}: {count}")

wb.close()
